# src/report.py
"""核查报告生成：PDF / Excel / 网页三种形式。

内容按《项目思路与实施流程》第八步的 13 项组织：
  1 项目基本信息  2 核查材料范围  3 数据提取结果  4 规则校验结果  5 异常问题清单
  6 原文证据及页码  7 差异计算过程  8 风险提示  9 修改建议  10 人工复核记录
  11 提取准确率与分项指标  12 CSI 公式/权重/分级  13 各类规则数量统计

★ 首页必须含免责声明（interfaces.md §6.3(3)）：
  报告只呈现核查结果，不作审计结论、不构成投资建议。
★ 报告中的一切数值与页码均取自程序产物，本模块不做任何计算与改写。
"""
from __future__ import annotations

import json
import time
from pathlib import Path

from decimal import Decimal, InvalidOperation

ROOT = Path(__file__).resolve().parents[1]

DISCLAIMER = "本系统输出为财务数据一致性核查结果，不构成投资建议，所有结论需人工复核。"
REPORT_VERSION = "1.0"
STATUS_ZH = {"PASS": "通过", "FAIL": "异常", "REVIEW": "需人工复核", "NA": "不适用"}
DIM_ZH = {"C内": "表内勾稽", "C间": "表间一致", "C期": "跨期衔接",
          "C附": "主表与附注", "C研": "研报引用"}
# 提取字段所属报表的中文名与展示顺序（供「按报表分类」渲染）
STATEMENT_ZH = {"BS": "资产负债表", "IS": "利润表", "CF": "现金流量表",
                "EQ": "所有者权益变动表", "NOTE": "财务报表附注",
                "MDNA": "管理层讨论与分析"}
STATEMENT_ORDER = ("BS", "IS", "CF", "EQ", "NOTE", "MDNA")

# 模型评价段落的固定免责声明（由程序渲染，不由模型产出，避免与报告正文重复）
COMMENT_DISCLAIMER = ("本段由模型基于程序提取与派生计算的数据生成，仅供参考，"
                      "不构成投资建议，请结合原始披露人工复核。")


def _csi_caveat(csi: dict) -> str:
    """CSI 口径提示：口径未裁定（provisional）或存在全「不适用」维度时必须显式标注。

    ★ 满分 ≠ 全部维度通过校验：某维度可执行规则数为 0 时，其权重被按比例分摊到
      其余维度，CSI 只反映"可执行维度"的得分，属**临时/不完整**结果，
      不得被读成整体一致性的最终结论。无可标注内容时返回空串。
    """
    csi = csi or {}
    na = list(csi.get("na_dimensions") or [])
    parts = []
    if csi.get("provisional"):
        parts.append(f"口径状态：临时（provisional，profile={csi.get('profile') or '默认'}），"
                     f"评分权重与容差档位待业务负责人裁定，非最终业务评分。")
    if na:
        parts.append(f"本 CSI 为「临时/不完整」结果：{'、'.join(na)} 共 {len(na)} 个维度无规则可执行"
                     f"（全部为「不适用」），其权重已按比例分摊至其余维度；"
                     f"CSI 仅反映可执行维度的得分，不等同于全部维度均通过校验。")
    return "　".join(parts)


def _csi_exec_note(csi: dict) -> str:
    """CSI 可执行维度数统计：显式给出「可执行维度 / 总维度」，避免满分被误读为全维度通过。

    只统计 csi["dimensions"] 里 rules_applicable > 0 的维度数，**不重算任何得分**。
    """
    dims = (csi or {}).get("dimensions") or {}
    total = len(dims)
    exe = sum(1 for v in dims.values() if (v or {}).get("rules_applicable"))
    txt = f"可执行维度 {exe} 个 / 共 {total} 个维度。"
    if total - exe:
        txt += (f"另有 {total - exe} 个维度无可执行规则，按「不适用」处理，"
                f"其权重已按比例分摊至其余维度。")
    return txt


def _calls_text(calls: dict) -> str:
    """把 `manifest.calls` 统计展开为可读文本，替代直接抛 JSON。

    只读取程序写入的既有字段（缺省按零），不臆造任何数值；供 PDF / HTML 的
    「项目基本信息」表使用，使评审能直接看懂在线/缓存/未命中的含义。
    """
    if not calls:
        return "无模型调用记录。"
    total = calls.get("total", 0)
    api_ok = calls.get("api_ok", 0)
    cache_hit = calls.get("cache_hit", 0)
    cache_miss = calls.get("cache_miss", 0)
    failed = calls.get("failed", 0)
    parts = [f"模型调用 {total} 次：在线成功 {api_ok} 次、命中本地缓存 {cache_hit} 次"]
    if cache_miss:
        parts.append(f"离线未命中缓存 {cache_miss} 次（相关语义任务已转人工复核）")
    if failed:
        parts.append(f"调用失败 {failed} 次")
    return "；".join(parts) + "。"


def _text_row_height(text: str, col_width: int, base: float = 15.0) -> float:
    """按列宽估算成段文本换行后所需行高（Excel 手动行高），避免长文只显示首行。

    中文按 2 个字符宽计；按 `\\n` 分段各自估算后累加；下限为单行高度。
    """
    if not text:
        return base
    lines = 0
    for seg in str(text).split("\n"):
        units = sum(2 if ord(ch) > 0x2E80 else 1 for ch in seg)
        lines += max(1, -(-units // max(1, col_width)))
    return base * max(1, lines)


def _fields_by_statement(fields: list[dict]) -> list[tuple[str, str, list[dict]]]:
    """把提取字段按所属报表分类 → [(statement, 中文名, 记录列表), ...]（按固定顺序）。

    statement 取自解析记录（BS/IS/CF/EQ/NOTE/MDNA）；缺失或未知的归入「其他」。
    """
    buckets: dict[str, list[dict]] = {}
    for r in fields or []:
        buckets.setdefault(str(r.get("statement") or "OTHER"), []).append(r)
    order = [k for k in STATEMENT_ORDER if k in buckets]
    order += [k for k in buckets if k not in STATEMENT_ORDER]
    return [(k, STATEMENT_ZH.get(k, "其他"), buckets[k]) for k in order]


def _analysis_groups(analysis: dict) -> list[tuple[str, str, list[dict]]]:
    """财务分析指标按 group 归类 → [(group, 中文名, 指标列表), ...]，空组不返回。"""
    from src.financial_analysis import GROUP_ZH, group_metrics
    from src.forecast import GROUP_ZH as FC_ZH
    labels = {**GROUP_ZH, **FC_ZH}
    grouped = group_metrics((analysis or {}).get("metrics") or [])
    order = list(GROUP_ZH) + list(FC_ZH)
    return [(g, labels.get(g, g), grouped[g]) for g in order if grouped.get(g)]


def _accuracy(model: dict) -> dict | None:
    """结构化提取准确率 E：以人工标准答案（data/gold_standard/*.json）为分母逐条比对。

    ★ 覆盖范围必须显式标注：金标准只覆盖特定公司/期间的关键字段，因此这里的 E 是
      「关键字段抽样准确率」，不等于字段字典全量准确率；同时给出金标准覆盖率
      （金标准条目数 ÷ 字段字典项数），避免把抽样 E 读成全量结论。
      找不到与本次运行匹配的金标准（关键字段重合不足）时返回 None，报告写明原因。
    """
    fields = model.get("fields") or []
    gold_dir = ROOT / "data" / "gold_standard"
    if not fields or not gold_dir.exists():
        return None
    keys = {(r.get("item_code"), r.get("period")) for r in fields}
    best = None
    for p in sorted(gold_dir.glob("*.json")):
        gold = _read(p)
        ents = (gold or {}).get("entries") or []
        if not ents:
            continue
        hit = sum(1 for e in ents if (e.get("item_code"), e.get("period")) in keys)
        cov = hit / len(ents)                     # 关键字段重合度
        # 重合不足视为"不是同一份材料"，不得拿来算分——否则是张冠李戴
        if cov < 0.5:
            continue
        if best is None or cov > best[2]:
            best = (p, gold, cov)
    if best is None:
        return None
    p, gold, cov = best
    from tools.check_accuracy import compare
    res = compare(gold, fields)
    try:
        import yaml
        dict_size = len(yaml.safe_load(
            (ROOT / "rules" / "fields.yaml").read_text(encoding="utf-8"))["fields"])
    except Exception:                                        # noqa: BLE001
        dict_size = None
    res["gold_file"] = p.name
    res["gold_entries"] = len(gold["entries"])
    res["match_ratio"] = round(cov, 4)
    res["dict_size"] = dict_size
    res["dict_coverage"] = (round(len(gold["entries"]) / dict_size, 4)
                            if dict_size else None)
    return res


def _accuracy_scope(acc: dict | None) -> list[str]:
    """提取准确率的**事实陈述**（数值 + 口径 + 覆盖范围），三种报告共用同一份文案。"""
    if not acc:
        return ["本次运行没有与之匹配的人工标准答案（data/gold_standard/），"
                "未计算提取准确率 E；该指标属人工抽样核查，不影响上述规则结论。"]
    cov = acc.get("dict_coverage")
    cov_txt = f"{cov * 100:.1f}%" if cov is not None else "—"
    return [
        f"结构化提取准确率 E = {acc.get('E')}%（正确 {acc.get('correct')} ÷ "
        f"应提取 {acc.get('total')}；错误 {acc.get('wrong')} 条、缺失 {acc.get('missing')} 条）。",
        f"口径与覆盖范围：真值取自人工标准答案《{acc.get('gold_file')}》"
        f"（{acc.get('gold_entries')} 条关键字段）；字段字典共 {acc.get('dict_size') or '—'} 项，"
        f"金标准覆盖率 {cov_txt}。",
        "★ 因此本 E 为「关键字段抽样准确率」，不等同于字段字典全量准确率；"
        "未列入金标准的字段需人工核对，本报告不对其做自动判定。",
    ]


def _commentary_sections(text: str) -> list[list[str]]:
    """把模型输出的「【小节名】正文」多行文本拆成 [[小节名, 正文], ...]。

    对齐券商点评研报的骨架：小节名固定为 核心观点/收入分析/盈利能力/现金流/风险提示（五小节）。
    不以「【】」开头的行并入上一小节；首行即无标题时按无标题正文处理。
    """
    out: list[list[str]] = []
    for raw in (text or "").splitlines():
        line = raw.strip()
        if not line:
            continue
        head, body = "", line
        if line.startswith("【") and "】" in line:
            end = line.index("】")
            head, body = line[1:end].strip(), line[end + 1:].strip()
        if head:
            out.append([head, body])
        elif out:
            out[-1][1] += body
        else:
            out.append(["", body])
    return out


def _read(p: Path, default=None):
    try:
        return json.loads(p.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return default


def collect(run_id: str) -> dict:
    """把一次运行的全部产物汇总成报告模型。缺哪份产物就少哪一节，不报错。"""
    base = ROOT / "runs" / run_id
    manifest = _read(base / "manifest.json", {}) or {}
    results = _read(base / "rules" / "results.json", {}) or {}
    csi = _read(base / "rules" / "csi.json", {}) or {}
    research = _read(base / "rules" / "research_claims.json", {}) or {}
    mappings = _read(base / "llm" / "mappings.json", {}) or {}
    commentary = _read(base / "llm" / "commentary.json", {}) or {}
    # 材料识别日志：哪些文件被当成什么处理、哪些没被识别，必须能呈现在报告上
    ingest_log = _read(base / "ingest_log.json", []) or []

    rule_results = results.get("results", [])
    # 模型解释单独存放（结果文件保持纯确定性），此处按编号关联回来仅供展示
    from src.explain import load_explanations
    explanations = load_explanations(run_id)
    for r in rule_results:
        e = explanations.get(r.get("rule_id"))
        if e:
            r["explanation"] = e
    for r in research.get("results", []):
        e = explanations.get(r.get("claim_id"))
        if e:
            r["explanation"] = e

    anomalies = [r for r in rule_results if r.get("status") in ("FAIL", "REVIEW")]
    research_fail = [r for r in research.get("results", []) if r.get("status") == "FAIL"]

    # 提取到的字段：现场解析目录优先，其次共享基线
    fields, ingested = [], base / "ingested"
    scan_dir = ingested if ingested.exists() else ROOT / "data" / "structured"
    for p in sorted(scan_dir.glob("*.json")) if scan_dir.exists() else []:
        payload = _read(p)
        if isinstance(payload, list):
            fields.extend(r for r in payload
                          if isinstance(r, dict) and "item_code" in r)

    model = {
        "run_id": run_id,
        "generated_at": time.strftime("%Y-%m-%dT%H:%M:%S+08:00"),
        "report_version": REPORT_VERSION,
        "disclaimer": DISCLAIMER,
        "manifest": manifest,
        "results": results,
        "rule_results": rule_results,
        "anomalies": anomalies,
        "csi": csi,
        "research": research,
        "research_fail": research_fail,
        "mappings": mappings,
        "fields": fields,
        "commentary": commentary,
        "ingest_log": ingest_log,
        "counts": _counts(rule_results, research),
    }
    # 成段结论：由已完成聚合的计数与比对结果生成，供三种报告共用同一份文案
    model["summary"] = build_summary(model)
    model["research_review"] = build_research_review(model)
    # 确定性财务分析（盈利能力/费用率/趋势/结构）：由提取字段直接计算，无模型参与
    from src.financial_analysis import analyze as analyze_financials
    from src.forecast import forecast as forecast_financials
    analysis = analyze_financials(fields)
    fc = forecast_financials(fields)
    # 盈利预测（情景测算，选题4基础）：并入「财务比率分析」章节末尾，与比率共用同一渲染路径
    analysis["metrics"] = (analysis.get("metrics") or []) + (fc.get("metrics") or [])
    analysis["warnings"] = (analysis.get("warnings") or []) + (fc.get("warnings") or [])
    model["analysis"] = analysis
    model["forecast"] = fc
    # 提取准确率 E：以人工标准答案为分母，无匹配金标准时为 None（报告写明原因）
    model["accuracy"] = _accuracy(model)
    return model


def _counts(rule_results: list[dict], research: dict) -> dict:
    """第 13 项：各类规则的适用/通过/异常/待复核数量。"""
    per_dim: dict[str, dict] = {}
    for r in rule_results:
        d = per_dim.setdefault(r.get("dimension", "?"), {
            "applicable": 0, "pass": 0, "fail": 0, "review": 0, "na": 0})
        st = r.get("status")
        if st == "NA":
            d["na"] += 1
            continue
        d["applicable"] += 1
        d[{"PASS": "pass", "FAIL": "fail", "REVIEW": "review"}.get(st, "review")] += 1
    rc = research.get("results", [])
    return {
        "rules_total": len(rule_results),
        "rules_per_dim": per_dim,
        "research_total": len(rc),
        "research_pass": sum(1 for r in rc if r.get("status") == "PASS"),
        "research_fail": sum(1 for r in rc if r.get("status") == "FAIL"),
        "research_review": len(research.get("review", [])),
        "fields": 0,      # 由 caller 填
    }


def merged_anomaly_counts(rule_results: list[dict], research: dict) -> tuple[int, int]:
    """「异常 / 待复核」的唯一口径：(异常数, 待复核数)。

    ★ 必须合并两个数据源，网页与三种报告都调这一个函数：
      - 规则级结果落在 runs/<run_id>/rules/results.json；
      - 研报引用比对结果落在 runs/<run_id>/rules/research_claims.json，
        其中 results 段是 PASS/FAIL，review 段全部是「程序无法自动判定」的
        待复核条目（status=REVIEW）。
      历史上网页只统计前者，于是显示 0，与报告总评的「不一致 3 条、
      待复核 41 条」对不上。计数口径必须单一来源，不得各算各的。
    """
    research = research or {}
    res = list(rule_results or []) + list(research.get("results") or [])
    # review 段本身就是"程序无法自动判定"的待复核清单，按条数计入即可；
    # 不要求每条都带 status（旧产物/测试造的数据可能没有该字段）。
    review = list(research.get("review") or [])
    n_fail = sum(1 for r in res if r.get("status") == "FAIL")
    n_review = sum(1 for r in res if r.get("status") == "REVIEW") + len(review)
    return n_fail, n_review


def _review_sentence(item: dict) -> str:
    """待复核条目的研报原句：优先 evidence.sentence，其次顶层 sentence（兼容旧产物）。"""
    return (item.get("evidence") or {}).get("sentence") or item.get("sentence") or ""


def research_review_rows(model: dict) -> list[list[str]]:
    """研报引用「待人工复核」条目的表格行：(科目, 期间, 原因, 研报原句)。

    ★ 只给计数不给明细，评审无从核对到底是哪几句没比对上；报告必须逐条列出。
    """
    return [[x.get("item_code") or "—", str(x.get("period") or "—"),
             x.get("reason", ""), _review_sentence(x)]
            for x in (model.get("research") or {}).get("review") or []]


def research_review_reasons(model: dict) -> list[str]:
    """待复核条目的原因分布（按出现次数降序）：「原因：N 条」。"""
    from collections import Counter
    c = Counter(x.get("reason", "") or "（未给出原因）"
                for x in (model.get("research") or {}).get("review") or [])
    return [f"{reason}：{n} 条" for reason, n in c.most_common()]


def ingest_notes(model: dict) -> list[str]:
    """材料识别结果的事实陈述：**未识别的文件必须显式列出**，不得静默丢弃。

    数据取自 runs/<run_id>/ingest_log.json（src.ingest.process_folder 的 log 落盘）。
    旧运行没有该文件时返回空列表，报告照常生成。
    """
    head = {"annual": "年报", "research": "研究报告",
            "unknown": "未识别", "error": "处理失败"}
    notes = []
    for e in model.get("ingest_log") or []:
        text = (f"[{head.get(e.get('kind'), e.get('kind'))}] "
                f"{e.get('file')}：{e.get('action')}")
        if e.get("kind") in ("unknown", "error"):
            text += ("——该文件未参与字段提取与一致性核查，如需核查请确认材料类型后重新上传。")
        notes.append(text)
    return notes


def _dim_counts(counts: dict) -> tuple[int, int, int, int, int]:
    """把 rules_per_dim 的计数做展示性汇总：(可执行, 通过, 异常, 待复核, 不适用)。

    只汇总已存在的计数器，不重算任何财务数值——报告不承担计算职责。
    """
    per = counts.get("rules_per_dim") or {}
    return (sum(v["applicable"] for v in per.values()),
            sum(v["pass"] for v in per.values()),
            sum(v["fail"] for v in per.values()),
            sum(v["review"] for v in per.values()),
            sum(v["na"] for v in per.values()))


def build_summary(model: dict) -> list[str]:
    """核查总评：把已有的计数与 CSI 组织成**成段结论**，置于报告首页。

    ★ 只陈述 model 里已经存在的数字（counts / csi / manifest / fields），
      不新做任何计算与猜测：报告不产结论性判断，仅呈现程序侧的事实。
      全部由确定性文本模板生成，离线可复现，不增加模型调用。
    """
    c = model.get("counts") or {}
    applicable, passed, failed, review, na = _dim_counts(c)
    paras = [f"本次核查共提取结构化字段 {len(model.get('fields') or [])} 条，"
             f"执行规则 {c.get('rules_total', 0)} 条，其中可执行 {applicable} 条、"
             f"通过 {passed} 条、异常 {failed} 条、待人工复核 {review} 条、不适用 {na} 条。"]

    per = c.get("rules_per_dim") or {}
    hits = [(k, v["fail"]) for k, v in sorted(per.items()) if v.get("fail")]
    if hits:
        paras.append("异常集中在 " + "、".join(
            f"{DIM_ZH.get(k, k)} {n} 条" for k, n in hits) + "。")
    else:
        paras.append("未发现规则级异常：表内勾稽与表间比对在容差内自洽。")

    rc = (model.get("research") or {}).get("results") or []
    if rc:
        paras.append(f"研报引用核查共比对 {len(rc)} 条："
                     f"一致 {sum(1 for r in rc if r.get('status') == 'PASS')} 条、"
                     f"不一致 {sum(1 for r in rc if r.get('status') == 'FAIL')} 条、"
                     f"待人工复核 {len((model.get('research') or {}).get('review') or [])} 条。")
    else:
        n_rv = len((model.get("research") or {}).get("review") or [])
        if n_rv:
            # ★ 有引用却一条都没能自动比对（离线缓存未命中、模型漏答）：不能说
            #   "未包含研究报告"，那是误导——必须如实标注这批引用转了人工复核。
            paras.append(f"研报引用核查本次未完成自动比对：{n_rv} 条引用全部转人工复核"
                         f"（模型未给出可用的科目/期间匹配，常见于离线运行或缓存未命中），"
                         f"研报引用维度本轮不参与 CSI 评分。")
        else:
            paras.append("本批材料未包含研究报告，或研报引用未匹配到可比字段："
                         "研报引用维度不适用。")

    csi = model.get("csi") or {}
    if csi.get("csi") is not None:
        paras.append(f"综合一致性指数 CSI = {csi.get('csi')}，"
                     f"等级 {csi.get('grade')}（{csi.get('grade_label', '')}）。")
    else:
        paras.append("本次运行没有可执行的规则，未计算 CSI。")

    advice = []
    if failed:
        advice.append(f"建议优先人工核对上述 {failed} 条异常，回到原始披露确认差异是否属于"
                      f"列报口径、会计政策变更或四舍五入。")
    if review:
        advice.append(f"另有 {review} 条标记为待人工复核，程序无法自动判定，不代表存在错误。")
    if not failed and not review:
        advice.append("报告期内没有需要人工介入的条目；可结合分部与附注披露做进一步复核。")
    calls = (model.get("manifest") or {}).get("calls") or {}
    if calls.get("cache_miss") or calls.get("failed"):
        advice.append(f"离线运行下有 {calls.get('cache_miss', 0)} 条语义任务未命中缓存、"
                      f"{calls.get('failed', 0)} 条调用失败，已转入人工复核，"
                      f"在线重跑可补齐这部分语义匹配。")
    paras.append("处置建议：" + "".join(advice))
    paras.append("本报告仅反映材料内部的一致性与引用关系，不构成审计意见或投资建议。")
    return paras


def build_research_review(model: dict) -> list[str]:
    """研报观点与财报数字差异点评：把已有比对结果写成成段点评。

    ★ 数值一律取自 research_check 的产出（left/right/diff/tolerance/cause），
      本函数只做文字陈述，不重算、不改写。
    """
    research = model.get("research") or {}
    rc = research.get("results") or []
    if not rc:
        return []
    fails = [r for r in rc if r.get("status") == "FAIL"]
    passed = [r for r in rc if r.get("status") == "PASS"]
    review = research.get("review") or []
    paras = [f"本批材料共比对研报引用 {len(rc)} 条：一致 {len(passed)} 条、"
             f"不一致 {len(fails)} 条、待人工复核 {len(review)} 条。"]
    for r in fails:
        sentence = (r.get("evidence") or {}).get("sentence") or "（研报原句未记录）"
        left = r.get("left") or {}
        right = r.get("right") or {}
        paras.append(
            f"研报称「{sentence}」——该句把 {r.get('item_code')}（{r.get('period')}）"
            f"表述为 {left.get('raw') or left.get('value')}，而财报披露为 {right.get('value')}，"
            f"差异 {r.get('diff')}（容差 {r.get('tolerance')}）；"
            f"归因 {r.get('cause', '')}：{r.get('reason', '') or '程序未给出进一步归因'}。"
            f"建议核对引用数值、单位与期间，并按证据定位回研报原文确认。")
    if not fails:
        paras.append("研报引用与财报数字全部一致：本批比对未发现引用层面的不一致。")
    if review:
        paras.append(f"上述 {len(review)} 条待人工复核条目的原因分布："
                     + "；".join(research_review_reasons(model)) + "。")
        paras.append("逐条明细（科目 / 期间 / 原因 / 研报原句）在同一章节的表格中列出；"
                     "待复核仅表示程序无法自动判定，不代表存在错误，"
                     "需按原句回到研报与财报原文核对后再定性。")
    return paras


def _suggest(status: str, dimension: str, reason: str) -> str:
    """第 9 项：修改建议。按状态与维度给通用建议，不替人下财务结论。"""
    if status == "REVIEW":
        return f"请人工核对{reason or '该条涉及的原始披露'}后回填复核结论。"
    if dimension == "C研":
        return "请核对研报引用的数值、单位与期间是否与财报一致，并按证据定位回原始材料确认。"
    if dimension == "C期":
        return "请确认是否存在会计政策变更或追溯重述；若无，则需核对上期期末数。"
    return "请按证据定位回原始报表核对相关科目，确认差异是否属于列报口径差异。"


# ---------------- PDF ----------------

def render_pdf(model: dict, path: Path) -> Path:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.platypus import (PageBreak, Paragraph, SimpleDocTemplate,
                                    Spacer, Table, TableStyle)

    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    base = getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=base["Title"], fontName="STSong-Light", fontSize=18)
    h2 = ParagraphStyle("h2", parent=base["Heading2"], fontName="STSong-Light", fontSize=13,
                        spaceBefore=10, spaceAfter=4)
    h3 = ParagraphStyle("h3", parent=base["Heading3"], fontName="STSong-Light", fontSize=10.5,
                        spaceBefore=6, spaceAfter=2, textColor=colors.HexColor("#2d6cdf"))
    body = ParagraphStyle("body", parent=base["BodyText"], fontName="STSong-Light",
                          fontSize=9.5, leading=14)
    small = ParagraphStyle("small", parent=body, fontSize=8, textColor=colors.grey)
    # 表格单元格样式：独立于 body，行高紧凑，配合 Paragraph 在列宽内自动换行
    cell = ParagraphStyle("cell", parent=base["BodyText"], fontName="STSong-Light",
                          fontSize=8, leading=10)

    def tbl(rows, widths=None, header=True):
        # ★ 单元格一律用 Paragraph 渲染：纯字符串不会被 reportlab 换行，
        #   长科目名（如"处置固定资产、无形资产和其他长期资产的损失…"）会直接
        #   溢出到相邻列、字体重叠。Paragraph 会在列宽内自动断行，杜绝重叠。
        rows = [[Paragraph(_esc(c), cell) if isinstance(c, str) else c for c in row]
                for row in rows]
        t = Table(rows, colWidths=widths, repeatRows=1 if header else 0)
        t.setStyle(TableStyle([
            ("FONTNAME", (0, 0), (-1, -1), "STSong-Light"),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#CCCCCC")),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F0F0F0")) if header else
            ("BACKGROUND", (0, 0), (-1, -1), colors.white),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ]))
        return t

    story = [Paragraph("财报一致性核查报告", h1),
             Paragraph(f"运行编号 {model['run_id']} ｜ 生成时间 {model['generated_at']}", small),
             Spacer(1, 4),
             Paragraph(f"<b>{DISCLAIMER}</b>", body),
             Spacer(1, 8)]

    # 核查总评：成段结论前置，让评审第一眼看到核查规模、异常分布与 CSI 等级
    story.append(Paragraph("核查总评", h2))
    for p in model.get("summary") or []:
        story.append(Paragraph(_esc(p), body))
    story.append(Spacer(1, 8))

    # 1 项目基本信息 / 2 材料范围
    m = model["manifest"]
    story.append(Paragraph("一、项目基本信息与核查材料范围", h2))
    story.append(tbl([
        ["项", "值"],
        ["运行编号", model["run_id"]],
        ["模型标识", (m.get("model") or {}).get("model_id", "—")],
        ["提示词哈希", "；".join(f"{p['prompt_id']}:{p['sha256'][:12]}"
                              for p in (m.get("prompts") or [])) or "—"],
        ["调用统计", _calls_text(m.get("calls", {}))],
        ["提取字段数", str(len(model["fields"]))],
        ["规则结果数", str(model["counts"]["rules_total"])],
    ], widths=[35 * mm, 130 * mm]))

    # 材料识别结果：未识别的文件必须显式标注，不得静默丢弃（否则用户以为它被核查了）
    notes = ingest_notes(model)
    if notes:
        story.append(Paragraph(f"材料识别结果（扫描 {len(model['ingest_log'])} 个文件）", h3))
        for t in notes:
            story.append(Paragraph("· " + _esc(t), small))
        story.append(Spacer(1, 6))

    # 3 数据提取结果（按报表类型分类展示）
    story.append(Paragraph("二、数据提取结果（按报表分类）", h2))
    if model["fields"]:
        for stmt, zh, recs in _fields_by_statement(model["fields"]):
            story.append(Paragraph(f"{zh}（{len(recs)} 条）", h3))
            rows = [["科目", "期间", "数值", "定位"]]
            for r in recs:
                loc = f"p{r['page']}" if r.get("page") is not None else (r.get("locator") or "—")
                rows.append([r.get("item_raw", r.get("item_code", "")), str(r.get("period", "")),
                             r.get("value", ""), loc])
            story.append(tbl(rows, widths=[46 * mm, 26 * mm, 48 * mm, 45 * mm]))
        story.append(Paragraph(
            f"本节按报表类型分类列出全部 {len(model['fields'])} 条提取记录（表格跨页自动续表）；"
            f"机读清单见 Excel 版「提取字段」页。", small))
    else:
        story.append(Paragraph("未提取到字段记录。", body))

    # 4 财务比率分析（程序确定性计算，无模型参与、可复现）
    story.append(PageBreak())
    story.append(Paragraph("三、财务比率分析", h2))
    analysis = model.get("analysis") or {}
    groups = _analysis_groups(analysis)
    if groups:
        for _g, gzh, metrics in groups:
            story.append(Paragraph(gzh, h3))
            rows = [["指标", "期间", "数值", "口径说明"]]
            for m in metrics:
                rows.append([m.get("name", ""), str(m.get("period", "")),
                             f"{m.get('value','')}{m.get('unit','')}", m.get("basis", "")])
            story.append(tbl(rows, widths=[36 * mm, 22 * mm, 26 * mm, 81 * mm]))
    else:
        story.append(Paragraph(
            "本次运行没有可计算财务比率的基础字段（如营业收入/营业成本/所有者权益等），未产出比率分析。",
            body))
    for w in analysis.get("warnings") or []:
        story.append(Paragraph("· " + _esc(w), small))
    if analysis.get("warnings"):
        story.append(Paragraph(
            "以上提示表示相应指标因输入缺失未予计算，不含推算或估计值。", small))

    # 5/6/7/8 异常清单（含证据与差异过程）
    story.append(PageBreak())
    story.append(Paragraph("四、异常问题清单（含证据与差异计算过程）", h2))
    if model["anomalies"]:
        for i, a in enumerate(model["anomalies"], 1):
            ev = a.get("evidence") or {}
            loc = f"p{ev['page']}" if ev.get("page") is not None else (ev.get("locator") or "—")
            story.append(Paragraph(
                f"<b>异常 {i}｜{a.get('rule_id')}｜{DIM_ZH.get(a.get('dimension',''), a.get('dimension',''))}"
                f"｜{STATUS_ZH.get(a.get('status'), a.get('status'))}</b>", body))
            story.append(Paragraph(
                f"计算式：{a.get('expr') or '—'}<br/>"
                f"差异：{a.get('diff')}　容差：{a.get('tolerance')}　"
                f"严重程度：{a.get('severity', '—')}<br/>"
                f"证据：{ev.get('file') or '—'}　定位：{loc}　片段：{ev.get('raw_text') or '—'}<br/>"
                f"建议：{_suggest(a.get('status'), a.get('dimension',''), a.get('reason',''))}", small))
            story.append(Spacer(1, 5))
    else:
        story.append(Paragraph("本次运行未发现异常。", body))

    # 研报引用异常与待复核：异常逐条给数值，待复核逐条给原因与原句
    _rv_rows = research_review_rows(model)
    if model["research_fail"] or _rv_rows:
        story.append(Paragraph("五、研报引用异常与待人工复核", h2))
        if model["research_fail"]:
            rows = [["科目", "期间", "研报值", "财报值", "差异", "归因"]]
            for r in model["research_fail"]:
                rows.append([r.get("item_code", ""), str(r.get("period", "")),
                             (r.get("left") or {}).get("raw", ""),
                             (r.get("right") or {}).get("value", "")[:18],
                             (r.get("diff") or "")[:18],
                             f"{r.get('cause','')} {r.get('reason','')}"])
            story.append(tbl(rows, widths=[26 * mm, 20 * mm, 30 * mm, 34 * mm,
                                           26 * mm, 29 * mm]))
        if _rv_rows:
            story.append(Paragraph(
                f"待人工复核 {len(_rv_rows)} 条（程序无法自动判定，不代表存在错误）："
                + "；".join(research_review_reasons(model)), small))
            rows = [["科目", "期间", "原因", "研报原句"]]
            rows.extend(_rv_rows)
            story.append(tbl(rows, widths=[24 * mm, 18 * mm, 62 * mm, 61 * mm]))
            story.append(Spacer(1, 5))

    # 研报观点 vs 财报数字：成段点评（有研报比对结果时才出现）
    review = model.get("research_review") or []
    if review:
        story.append(Paragraph("研报观点与财报数字差异点评", h2))
        for p in review:
            story.append(Paragraph(_esc(p), body))
        story.append(Spacer(1, 6))

    # 大模型研报式评价（仅供参考）：数字均来自程序提取与派生计算，经数字白名单校验
    comment = (model.get("commentary") or {}).get("text", "")
    if comment:
        story.append(Paragraph("大模型经营评价（仅供参考）", h2))
        for sec_head, sec_text in _commentary_sections(comment):
            if sec_head:
                story.append(Paragraph(_esc(sec_head), h3))
            if sec_text:
                story.append(Paragraph(_esc(sec_text), body))
        story.append(Paragraph(COMMENT_DISCLAIMER, small))
        story.append(Spacer(1, 6))

    # 11/12/13 量化评价
    story.append(PageBreak())
    story.append(Paragraph("六、一致性量化评价", h2))
    csi = model["csi"]
    if csi:
        rows = [["维度", "含义", "可执行", "通过", "异常", "待复核", "不适用", "得分"]]
        for k, v in sorted((csi.get("dimensions") or {}).items()):
            rows.append([k, DIM_ZH.get(k, ""), v["rules_applicable"], v["pass"], v["fail"],
                         v["review"], v["na"],
                         "不适用" if v["score"] is None else f"{v['score']*100:.1f}"])
        story.append(tbl(rows, widths=[16 * mm, 30 * mm, 18 * mm, 14 * mm, 14 * mm,
                                       18 * mm, 16 * mm, 18 * mm]))
        story.append(Spacer(1, 4))
        story.append(Paragraph(
            f"<b>综合一致性指数 CSI = {csi.get('csi')}　等级 {csi.get('grade')}"
            f"（{csi.get('grade_label','')}）</b>", body))
        story.append(Paragraph(
            f"计算公式：CSI = Σ(维度得分 × 归一化权重) × 100；"
            f"维度得分 = (通过数 + 0.5 × 待复核数) ÷ 可执行规则数。<br/>"
            f"原权重：{json.dumps(csi.get('weights_used', {}), ensure_ascii=False)}<br/>"
            f"归一化权重：{json.dumps(csi.get('weights_normalized', {}), ensure_ascii=False)}<br/>"
            f"不适用维度：{','.join(csi.get('na_dimensions') or []) or '无'}　"
            f"（{csi.get('note','')}）", small))
        # 可执行维度数统计：显式标注可执行维度 / 总维度，避免满分被误读为全维度通过
        story.append(Spacer(1, 3))
        story.append(Paragraph(f"<b>【可执行维度】</b>{_esc(_csi_exec_note(csi))}", small))
        caveat = _csi_caveat(csi)
        if caveat:
            story.append(Spacer(1, 3))
            story.append(Paragraph(f"<b>【口径提示】</b>{_esc(caveat)}", small))
    else:
        story.append(Paragraph("本次运行无可执行的规则，未计算 CSI。", body))

    # 11 提取准确率（E）：以人工标准答案为分母，并如实标注覆盖范围
    acc = model.get("accuracy")
    story.append(Paragraph("七、提取准确率（E）", h2))
    for t in _accuracy_scope(acc):
        story.append(Paragraph(_esc(t), body if t.startswith("结构化") else small))
    if acc:
        bad = [d for d in (acc.get("details") or []) if d["判定"] != "正确"][:15]
        if bad:
            rows = [["判定", "科目", "期间", "标准值", "提取值"]]
            for d in bad:
                rows.append([d["判定"], d["item_code"], str(d["period"]),
                             d["标准值"], str(d["提取值"])])
            story.append(tbl(rows, widths=[16 * mm, 34 * mm, 26 * mm, 45 * mm, 43 * mm]))
        else:
            story.append(Paragraph("列入金标准的关键字段全部提取正确。", small))
    story.append(Spacer(1, 6))

    # 8 风险提示 / 9 修改建议 / 10 人工复核
    story.append(Paragraph("八、风险提示与人工复核", h2))
    story.append(Paragraph(
        "1. 本报告仅反映材料内部的一致性与引用关系，<b>不构成审计意见</b>。<br/>"
        "2. 标记为「需人工复核」的条目表示程序无法确定，<b>不代表存在错误</b>。<br/>"
        "3. 差异可能源于列报口径、会计政策变更或四舍五入，须结合原始披露判断。<br/>"
        "4. 人工复核结论请在结果文件的 human_review 字段回填，报告将同时呈现机器初判与人工结论。",
        small))
    story.append(Spacer(1, 6))
    story.append(Paragraph("人工复核记录", h2))
    reviewed = [r for r in model["rule_results"] if r.get("human_review")]
    # ★ 口径与网页指标卡、总评统一：规则级与研报两个来源合并，避免出现"待复核 0 条"
    #   却在上文列出 41 条待复核明细的自相矛盾。
    _n_fail, _n_review = merged_anomaly_counts(model.get("rule_results") or [],
                                               model.get("research") or {})
    story.append(Paragraph(
        f"已回填复核结论 {len(reviewed)} 条；异常 {_n_fail} 条、待人工复核 {_n_review} 条"
        f"（其中研报引用 {model['counts'].get('research_review', 0)} 条）。", small))
    story.append(Spacer(1, 10))
    story.append(Paragraph(DISCLAIMER, small))

    path.parent.mkdir(parents=True, exist_ok=True)
    SimpleDocTemplate(str(path), pagesize=A4,
                      leftMargin=18 * mm, rightMargin=18 * mm,
                      topMargin=16 * mm, bottomMargin=16 * mm,
                      title=f"财报一致性核查报告 {model['run_id']}").build(story)
    return path


# ---------------- Excel ----------------

def render_xlsx(model: dict, path: Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    head_font = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="F0F0F0")

    def sheet(name, header, rows):
        ws = wb.create_sheet(name)
        ws.append(header)
        for c in ws[1]:
            c.font, c.fill = head_font, head_fill
        for r in rows:
            ws.append(r)
        for i, _ in enumerate(header, 1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = 20
        ws.freeze_panes = "A2"
        return ws

    # 摘要
    ws = wb.active
    ws.title = "摘要"
    csi = model["csi"] or {}
    ws.append(["项", "值"])
    for c in ws[1]:
        c.font, c.fill = head_font, head_fill
    for k, v in [
        ("运行编号", model["run_id"]), ("生成时间", model["generated_at"]),
        ("模型标识", (model["manifest"].get("model") or {}).get("model_id", "—")),
        ("提取字段数", len(model["fields"])),
        ("规则结果数", model["counts"]["rules_total"]),
        ("CSI", csi.get("csi")), ("等级", f"{csi.get('grade','')} {csi.get('grade_label','')}"),
    ]:
        ws.append([k, v])
    caveat = _csi_caveat(csi)
    if caveat:
        ws.append(["CSI 口径提示", caveat])
        _cell = ws.cell(ws.max_row, 2)
        _cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[ws.max_row].height = _text_row_height(_cell.value, 70)
    if csi:
        ws.append(["CSI 可执行维度", _csi_exec_note(csi)])
        _cell = ws.cell(ws.max_row, 2)
        _cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[ws.max_row].height = _text_row_height(_cell.value, 70)
    # 提取准确率：与 PDF/网页一致，无论有无金标准都给出该行（无金标准时写明原因）
    acc = model.get("accuracy")
    if acc:
        cov = acc.get("dict_coverage")
        cov_txt = f"{cov * 100:.1f}%" if cov is not None else "—"
        acc_txt = (f"{acc.get('E')}%（正确 {acc.get('correct')}/{acc.get('total')}，"
                   f"金标准 {acc.get('gold_entries')} 条，字段字典覆盖率 {cov_txt}）"
                   f"——属关键字段抽样准确率，非字段字典全量准确率")
    else:
        acc_txt = "本次运行没有匹配的人工标准答案（data/gold_standard/），未计算提取准确率 E"
    ws.append(["提取准确率", acc_txt])
    ws.cell(ws.max_row, 2).alignment = Alignment(wrap_text=True, vertical="top")
    ws.append(["声明", DISCLAIMER])
    # 材料识别结果：未识别的文件在 Excel 里也必须能看到，不得只在命令行 stdout 一闪而过
    _notes = ingest_notes(model)
    if _notes:
        ws.append(["材料识别", "\n".join(_notes)])
        _cell = ws.cell(ws.max_row, 2)
        _cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[ws.max_row].height = _text_row_height(_cell.value, 70)
    # 核查总评 / 研报差异点评 / 大模型评价：成段文本整段落入单元格并自动换行，Excel 里也能直接读到结论
    comment = (model.get("commentary") or {}).get("text", "")
    for label, paras in (("核查总评", model.get("summary") or []),
                         ("研报差异点评", model.get("research_review") or []),
                         ("大模型经营评价(仅供参考)", [comment] if comment else [])):
        if not paras:
            continue
        ws.append([label, "\n".join(paras)])
        _cell = ws.cell(ws.max_row, 2)
        _cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[ws.max_row].height = _text_row_height(_cell.value, 70)
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 70

    # 异常清单
    sheet("异常清单", ["规则编号", "维度", "状态", "期间", "计算式", "差异", "容差",
                       "严重程度", "文件", "页码/定位", "原文片段", "修改建议"],
          [[r.get("rule_id"), DIM_ZH.get(r.get("dimension", ""), r.get("dimension")),
            STATUS_ZH.get(r.get("status"), r.get("status")), r.get("period", ""),
            r.get("expr", ""), r.get("diff"), r.get("tolerance"), r.get("severity"),
            (r.get("evidence") or {}).get("file"),
            (r.get("evidence") or {}).get("page")
            if (r.get("evidence") or {}).get("page") is not None
            else (r.get("evidence") or {}).get("locator"),
            (r.get("evidence") or {}).get("raw_text"),
            _suggest(r.get("status"), r.get("dimension", ""), r.get("reason", ""))]
           for r in model["anomalies"]])

    # 规则明细
    sheet("规则明细", ["规则编号", "维度", "状态", "期间", "期别", "差异", "容差", "表达式"],
          [[r.get("rule_id"), r.get("dimension"), r.get("status"), r.get("period", ""),
            r.get("period_type", ""), r.get("diff"), r.get("tolerance"), r.get("expr", "")]
           for r in model["rule_results"]])

    # 提取字段
    sheet("提取字段", ["科目编号", "原文科目", "期间", "期别", "数值", "单位",
                       "页码/定位", "来源片段"],
          [[r.get("item_code"), r.get("item_raw"), r.get("period"), r.get("period_type"),
            r.get("value"), r.get("unit_raw"),
            r.get("page") if r.get("page") is not None else r.get("locator"),
            r.get("raw_text")] for r in model["fields"]])

    # 财务比率分析（程序确定性计算，无模型参与）：新增独立表页
    yuan = [["分组", "指标", "期间", "数值", "单位", "口径说明"]]
    for _g, gzh, metrics in _analysis_groups(model.get("analysis") or {}):
        for m in metrics:
            yuan.append([gzh, m.get("name", ""), str(m.get("period", "")),
                         m.get("value", ""), m.get("unit", ""), m.get("basis", "")])
    ws = sheet("财务比率分析", ["分组", "指标", "期间", "数值", "单位", "口径说明"],
               yuan[1:])
    for w in (model.get("analysis") or {}).get("warnings") or []:
        ws.append(["提示", w, "", "", "", ""])
    ws.column_dimensions["F"].width = 60

    # 提取准确率明细（以人工标准答案逐条比对）
    if model.get("accuracy"):
        sheet("提取准确率", ["判定", "科目", "原文科目", "期间", "标准值", "提取值", "核对依据"],
              [[d["判定"], d["item_code"], d["item_raw"], d["period"],
                d["标准值"], d["提取值"], d["核对依据"]]
               for d in model["accuracy"]["details"]])

    # 研报引用
    rc = (model["research"] or {}).get("results", [])
    if rc:
        sheet("研报引用", ["状态", "科目", "期间", "研报原文值", "研报数值", "财报值",
                           "差异", "容差", "归因", "说明", "研报原句"],
              [[r.get("status"), r.get("item_code"), r.get("period"),
                (r.get("left") or {}).get("raw"), (r.get("left") or {}).get("value"),
                (r.get("right") or {}).get("value"), r.get("diff"), r.get("tolerance"),
                r.get("cause"), r.get("reason"), (r.get("evidence") or {}).get("sentence")]
               for r in rc])

    # 待复核
    sheet("待人工复核", ["类型", "编号/科目", "原因", "原文"],
          [["规则", r.get("rule_id"), r.get("reason", ""),
            (r.get("evidence") or {}).get("raw_text", "")] for r in model["rule_results"]
           if r.get("status") == "REVIEW"]
          + [["研报", x.get("claim_id", ""), x.get("reason", ""), _review_sentence(x)]
             for x in (model["research"] or {}).get("review", [])])

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


# ---------------- HTML ----------------

_HTML = """<!doctype html><html lang="zh-CN"><meta charset="utf-8">
<title>财报一致性核查报告 {run_id}</title>
<style>
 body{{font-family:"Microsoft YaHei",system-ui,sans-serif;margin:0;background:#f7f8fa;color:#222}}
 .wrap{{max-width:1100px;margin:0 auto;padding:24px}}
 h1{{font-size:22px;margin:0 0 4px}} h2{{font-size:16px;margin:28px 0 8px;
   border-left:4px solid #2d6cdf;padding-left:8px}}
 h3{{font-size:14px;margin:14px 0 4px;color:#2d6cdf}}
 .meta{{color:#666;font-size:13px}} .warn{{background:#fff4e5;border:1px solid #ffd08a;
   padding:10px 12px;border-radius:6px;font-size:13px;margin:12px 0}}
 table{{border-collapse:collapse;width:100%;background:#fff;font-size:13px}}
 th,td{{border:1px solid #e3e6ea;padding:6px 8px;text-align:left;vertical-align:top}}
 th{{background:#f0f2f5}} .FAIL{{color:#c0392b;font-weight:600}}
 .REVIEW{{color:#b7791f}} .PASS{{color:#1e7a46}} .NA{{color:#888}}
 .csi{{font-size:18px;font-weight:700;margin:8px 0}}
 code{{background:#f0f2f5;padding:1px 4px;border-radius:3px;font-size:12px}}
</style>
<div class="wrap">
<h1>财报一致性核查报告</h1>
<div class="meta">运行编号 {run_id} ｜ 生成时间 {generated_at} ｜ 报告版本 {report_version}</div>
<div class="warn">{disclaimer}</div>
{sections}
</div></html>"""


def _esc(s) -> str:
    return (str(s).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"))


def _table(header: list[str], rows: list[list]) -> str:
    out = ["<table><thead><tr>" + "".join(f"<th>{_esc(h)}</th>" for h in header) + "</tr></thead><tbody>"]
    for r in rows:
        cells = []
        for i, v in enumerate(r):
            cls = f' class="{_esc(v)}"' if i == 0 and str(v) in STATUS_ZH else ""
            cells.append(f"<td{cls}>{_esc(v)}</td>")
        out.append("<tr>" + "".join(cells) + "</tr>")
    return "".join(out) + "</tbody></table>"


def render_html(model: dict, path: Path) -> Path:
    sections = []
    m = model["manifest"]
    # 核查总评：成段结论置顶，便于评审直接读到核查规模、异常分布与 CSI 等级
    sections.append("<h2>核查总评</h2>" + "".join(
        f"<p>{_esc(p)}</p>" for p in (model.get("summary") or [])))
    sections.append("<h2>一、项目基本信息与材料范围</h2>" + _table(
        ["项", "值"],
        [["运行编号", model["run_id"]],
         ["模型标识", (m.get("model") or {}).get("model_id", "—")],
         ["提示词哈希", "；".join(f"{p['prompt_id']}:{p['sha256'][:12]}"
                              for p in (m.get("prompts") or [])) or "—"],
         ["调用统计", _calls_text(m.get("calls", {}))],
         ["提取字段数", len(model["fields"])],
         ["规则结果数", model["counts"]["rules_total"]]]))
    # 材料识别结果：未识别的文件必须显式标注，不得静默丢弃
    _notes = ingest_notes(model)
    if _notes:
        sections.append("<h3>材料识别结果</h3>" + "".join(
            f'<p class="meta">{_esc(t)}</p>' for t in _notes))

    # 数据提取结果（按报表类型分类展示）
    sections.append("<h2>二、数据提取结果（按报表分类）</h2>")
    if model["fields"]:
        for _stmt, zh, recs in _fields_by_statement(model["fields"]):
            sections.append(f"<h3>{_esc(zh)}（{len(recs)} 条）</h3>" + _table(
                ["科目", "期间", "数值", "定位"],
                [[r.get("item_raw", r.get("item_code", "")), str(r.get("period", "")),
                  r.get("value", ""),
                  f"p{r['page']}" if r.get("page") is not None else (r.get("locator") or "—")]
                 for r in recs]))
    else:
        sections.append("<p>未提取到字段记录。</p>")

    # 财务比率分析（程序确定性计算，无模型参与、可复现）
    analysis = model.get("analysis") or {}
    sections.append("<h2>三、财务比率分析</h2>")
    if _analysis_groups(analysis):
        for _g, gzh, metrics in _analysis_groups(analysis):
            sections.append(f"<h3>{_esc(gzh)}</h3>" + _table(
                ["指标", "期间", "数值", "口径说明"],
                [[m.get("name", ""), str(m.get("period", "")),
                  f"{m.get('value','')}{m.get('unit','')}", m.get("basis", "")]
                 for m in metrics]))
    else:
        sections.append("<p>本次运行没有可计算财务比率的基础字段（如营业收入/营业成本/"
                        "所有者权益等），未产出比率分析。</p>")
    for w in analysis.get("warnings") or []:
        sections.append(f'<p class="meta">{_esc(w)}</p>')

    sections.append("<h2>四、异常问题清单（含证据与差异过程）</h2>")
    if model["anomalies"]:
        sections.append(_table(
            ["状态", "规则", "维度", "计算式", "差异", "容差", "证据定位", "原文片段", "修改建议"],
            [[STATUS_ZH.get(r.get("status"), r.get("status")), r.get("rule_id"),
              DIM_ZH.get(r.get("dimension", ""), r.get("dimension")),
              r.get("expr", ""), r.get("diff"), r.get("tolerance"),
              f"p{(r.get('evidence') or {}).get('page')}"
              if (r.get("evidence") or {}).get("page") is not None
              else ((r.get("evidence") or {}).get("locator") or "—"),
              (r.get("evidence") or {}).get("raw_text") or "—",
              _suggest(r.get("status"), r.get("dimension", ""), r.get("reason", ""))]
             for r in model["anomalies"]]))
    else:
        sections.append("<p>本次运行未发现异常。</p>")

    # 研报引用异常与待复核：异常逐条给出数值与归因，待复核逐条给出原因与原句
    _rv_rows = research_review_rows(model)
    if model["research_fail"] or _rv_rows:
        sections.append("<h2>五、研报引用异常与待人工复核</h2>")
        if model["research_fail"]:
            sections.append(_table(
                ["状态", "科目", "期间", "研报原文", "研报值", "财报值", "差异", "归因"],
                [[r.get("status"), r.get("item_code"), r.get("period"),
                  (r.get("left") or {}).get("raw"), (r.get("left") or {}).get("value"),
                  (r.get("right") or {}).get("value"), r.get("diff"),
                  f"{r.get('cause','')} {r.get('reason','')}"]
                 for r in model["research_fail"]]))
        if _rv_rows:
            sections.append(f'<p class="meta">待人工复核 {len(_rv_rows)} 条'
                            f'（程序无法自动判定，不代表存在错误）：'
                            f'{_esc("；".join(research_review_reasons(model)))}</p>')
            sections.append(_table(["科目", "期间", "原因", "研报原句"], _rv_rows))

    # 研报观点 vs 财报数字：成段点评（有研报比对结果时才出现）
    review = model.get("research_review") or []
    if review:
        sections.append("<h2>研报观点与财报数字差异点评</h2>" + "".join(
            f"<p>{_esc(p)}</p>" for p in review))

    # 大模型研报式评价（仅供参考）
    comment = (model.get("commentary") or {}).get("text", "")
    if comment:
        parts = ["<h2>大模型经营评价（仅供参考）</h2>"]
        for sec_head, sec_text in _commentary_sections(comment):
            if sec_head:
                parts.append(f"<h3>{_esc(sec_head)}</h3>")
            if sec_text:
                parts.append(f"<p>{_esc(sec_text)}</p>")
        parts.append(f'<p class="warn">{_esc(COMMENT_DISCLAIMER)}</p>')
        sections.append("".join(parts))

    csi = model["csi"] or {}
    if csi:
        sections.append("<h2>六、一致性量化评价</h2>")
        sections.append(f'<div class="csi">CSI = {csi.get("csi")}　'
                        f'等级 {csi.get("grade")}（{csi.get("grade_label","")}）</div>')
        sections.append(_table(
            ["维度", "含义", "可执行", "通过", "异常", "待复核", "不适用", "得分"],
            [[k, DIM_ZH.get(k, ""), v["rules_applicable"], v["pass"], v["fail"], v["review"],
              v["na"], "不适用" if v["score"] is None else f"{v['score']*100:.1f}"]
             for k, v in sorted((csi.get("dimensions") or {}).items())]))
        sections.append(
            f'<p><code>CSI = Σ(维度得分 × 归一化权重) × 100</code>　'
            f'维度得分 = (通过 + 0.5×待复核) ÷ 可执行规则数<br/>'
            f'原权重 {_esc(json.dumps(csi.get("weights_used", {}), ensure_ascii=False))}　'
            f'归一化 {_esc(json.dumps(csi.get("weights_normalized", {}), ensure_ascii=False))}<br/>'
            f'不适用维度：{_esc(",".join(csi.get("na_dimensions") or []) or "无")}</p>')
        sections.append(f'<p><b>【可执行维度】</b>{_esc(_csi_exec_note(csi))}</p>')
        caveat = _csi_caveat(csi)
        if caveat:
            sections.append(f'<p class="warn"><b>【口径提示】</b>{_esc(caveat)}</p>')

    # 提取准确率（E）：以人工标准答案为分母，并如实标注覆盖范围
    acc = model.get("accuracy")
    sections.append("<h2>七、提取准确率（E）</h2>")
    for i, t in enumerate(_accuracy_scope(acc)):
        cls = ' class="csi"' if i == 0 else ""
        sections.append(f"<p{cls}>{_esc(t)}</p>")
    if acc:
        bad = [d for d in (acc.get("details") or []) if d["判定"] != "正确"][:15]
        if bad:
            sections.append(_table(["判定", "科目", "期间", "标准值", "提取值"],
                                   [[d["判定"], d["item_code"], str(d["period"]),
                                     d["标准值"], str(d["提取值"])] for d in bad]))
        else:
            sections.append("<p>列入金标准的关键字段全部提取正确。</p>")

    sections.append("<h2>八、风险提示与人工复核</h2>"
                    "<p>1. 本报告仅反映材料内部一致性与引用关系，<b>不构成审计意见</b>。<br/>"
                    "2. 「需人工复核」表示程序无法确定，<b>不代表存在错误</b>。<br/>"
                    "3. 差异可能源于列报口径、会计政策变更或四舍五入，须结合原始披露判断。</p>")

    html = _HTML.format(run_id=_esc(model["run_id"]),
                        generated_at=_esc(model["generated_at"]),
                        report_version=_esc(model["report_version"]),
                        disclaimer=_esc(DISCLAIMER),
                        sections="".join(sections))
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(html, encoding="utf-8")
    return path


def build_all(run_id: str, formats=("pdf", "xlsx", "html")) -> dict:
    """生成三种形式的报告，返回 {格式: 路径}。"""
    model = collect(run_id)
    model["counts"]["fields"] = len(model["fields"])
    out = ROOT / "runs" / run_id
    made = {}
    if "pdf" in formats:
        made["pdf"] = render_pdf(model, out / "report.pdf")
    if "xlsx" in formats:
        made["xlsx"] = render_xlsx(model, out / "report.xlsx")
    if "html" in formats:
        made["html"] = render_html(model, out / "report.html")
    return {k: str(v.relative_to(ROOT)) for k, v in made.items()}
